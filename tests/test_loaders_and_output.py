"""Folder scanning, document loading and .docx output."""

import tempfile
from pathlib import Path

from docx import Document

from core.doc_writer import write_combined_docx, write_summary_docx
from core.loaders import load_document, scan_folder


def _make_folder(root: Path) -> None:
    (root / "a.txt").write_text("plain text", encoding="utf-8")
    (root / "b.md").write_text("# markdown", encoding="utf-8")
    (root / "image.png").write_bytes(b"\x89PNG")
    (root / "sub").mkdir()
    (root / "sub" / "c.txt").write_text("nested", encoding="utf-8")
    doc = Document()
    doc.add_paragraph("first paragraph")
    doc.add_paragraph("   ")
    doc.add_paragraph("second paragraph")
    doc.save(str(root / "d.docx"))


def test_scan_folder_finds_only_supported_files():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        _make_folder(root)
        names = [p.name for p in scan_folder(tmp)]
        assert names == ["a.txt", "b.md", "d.docx", "c.txt"]
        assert [p.name for p in scan_folder(tmp, recursive=False)] == ["a.txt", "b.md", "d.docx"]


def test_scan_folder_rejects_missing_folder():
    try:
        scan_folder("/does/not/exist")
    except FileNotFoundError:
        return
    raise AssertionError("expected FileNotFoundError")


def test_load_document_reads_text_and_docx():
    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        _make_folder(root)
        assert load_document(root / "a.txt") == "plain text"
        assert load_document(root / "d.docx") == "first paragraph\n\nsecond paragraph"


def test_load_document_rejects_unsupported_type():
    try:
        load_document(Path("file.xlsx"))
    except ValueError:
        return
    raise AssertionError("expected ValueError")


def test_docx_writers_produce_readable_files():
    with tempfile.TemporaryDirectory() as tmp:
        single = Path(tmp) / "single.docx"
        write_summary_docx(single, "Title", "para one\n\npara two")
        texts = [p.text for p in Document(str(single)).paragraphs]
        assert texts == ["Title", "para one", "para two"]

        combined = Path(tmp) / "combined.docx"
        write_combined_docx(combined, "/data", [("a.txt", "sum a"), ("b.txt", "sum b")])
        texts = [p.text for p in Document(str(combined)).paragraphs if p.text]
        assert texts[0] == "Folder Summary: /data"
        assert "a.txt" in texts and "sum a" in texts and "b.txt" in texts and "sum b" in texts
