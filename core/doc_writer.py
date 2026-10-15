"""
core/doc_writer.py
===================
Helper functions that write summaries to disk as .docx (Word) files.
Uses the python-docx library (already in requirements.txt, also used
in core/loaders.py for READING docx files).

This module ONLY handles writing; api/main.py decides when/where to
write (based on the save_mode/save_format parameters).
"""

from pathlib import Path

from docx import Document as DocxDocument


def write_summary_docx(out_path: Path, title: str, summary: str) -> None:
    """Saves a single document's summary as a .docx file.
    The summary text is split on blank lines and each resulting piece
    is added as its own Word paragraph (rather than a single block)."""
    doc = DocxDocument()
    doc.add_heading(title, level=1)
    for paragraph in summary.split("\n\n"):
        paragraph = paragraph.strip()
        if paragraph:
            doc.add_paragraph(paragraph)
    doc.save(str(out_path))


def write_combined_docx(out_path: Path, folder: str, parts: list[tuple[str, str]]) -> None:
    """Merges the summaries of multiple documents into a SINGLE .docx file.
    parts: [(file_name, summary_text), ...] -- each under its own heading,
    separated by a page break."""
    doc = DocxDocument()
    doc.add_heading(f"Folder Summary: {folder}", level=1)

    for i, (file_name, summary) in enumerate(parts):
        doc.add_heading(file_name, level=2)
        for paragraph in summary.split("\n\n"):
            paragraph = paragraph.strip()
            if paragraph:
                doc.add_paragraph(paragraph)
        if i < len(parts) - 1:
            doc.add_page_break()

    doc.save(str(out_path))
